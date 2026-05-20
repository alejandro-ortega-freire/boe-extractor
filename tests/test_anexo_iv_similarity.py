import unittest
from tempfile import TemporaryDirectory
from pathlib import Path

from docx import Document

from source.anexo_iv_similarity import compare_mappings, mapping_from_anexo


class AnexoIVSimilarityTests(unittest.TestCase):
    def test_ra_and_c_are_treated_as_equivalent_for_truth_comparison(self):
        expected = [{"criterion": "C1", "numbers": {1, 2}}]
        actual = [{"criterion": "RA1", "numbers": {1, 2}}]

        score, details = compare_mappings(expected, actual)

        self.assertEqual(score, 1.0)
        self.assertEqual(details[0]["score"], 1.0)

    def test_truth_files_can_reuse_previous_content_when_original_has_empty_cell(self):
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "truth.docx"
            doc = Document()
            table = doc.add_table(rows=3, cols=2)
            table.cell(0, 0).text = "Objetivos"
            table.cell(0, 1).text = "Contenidos"
            table.cell(1, 0).text = "RA1"
            table.cell(1, 1).text = "1. Contenido compartido"
            table.cell(2, 0).text = "RA2"
            table.cell(2, 1).text = ""
            doc.save(path)

            expected = mapping_from_anexo(path, fill_empty_with_previous=True)

        self.assertEqual(
            expected,
            [
                {"criterion": "C1", "numbers": {1}},
                {"criterion": "C2", "numbers": {1}},
            ],
        )

    def test_optional_truth_documents_are_readable_when_available(self):
        from source.anexo_iv_similarity import TRUTH_FILES

        available = [path for path in TRUTH_FILES.values() if path.exists()]

        if not available:
            self.skipTest("No hay anexos IV verdad disponibles en esta máquina.")

        for path in available:
            with self.subTest(path=str(path)):
                self.assertIsInstance(mapping_from_anexo(path), list)
